from datetime import date
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook

from prod_guard import load_allowlist
from prod_smoke import ALLOWLIST_PATH, CreatedLog, smoke_label, write_prod_results

HQ = {"connectqa-automation-prod", "qateam", "casesearch"}


def test_allowlist_yaml_matches_spec():
    allowlist = load_allowlist(ALLOWLIST_PATH)
    assert allowlist.hosts["connect.dimagi.com"] == {
        "pm_automation_01", "network-manager", "program-manager-prod", "network-manager-prod",
    }
    assert allowlist.hosts["www.commcarehq.org"] == HQ
    assert allowlist.hosts["commcarehq.org"] == HQ


def test_smoke_label():
    assert smoke_label("CCCT-1234", date(2026, 10, 27)) == "QA-SMOKE CCCT-1234 27-Oct"


def test_created_log_leftovers():
    log = CreatedLog()
    log.add("opportunity", "pm_automation_01", "QA-SMOKE A")
    log.add("payment unit", "pm_automation_01", "QA-SMOKE B", url="https://x/")
    log.mark_deleted("QA-SMOKE A")
    assert log.leftovers() == [
        {"kind": "payment unit", "org": "pm_automation_01", "name": "QA-SMOKE B", "url": "https://x/"}
    ]


def _plan(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Feature"
    ws.append(["Test Case ID", "Description", "Steps", "Expected Outcome", "Status",
               "Comments / Notes", "Automatable"])
    ws.append(["TC-FT-001", "one", "s", "e", "Pass", "", "Yes"])
    ws.append(["TC-FT-002", "two", "s", "e", None, "", "Yes"])
    path = tmp_path / "CCCT-1_Feature_Test_Plan.xlsx"
    wb.save(path)
    return path


def _runner():
    return SimpleNamespace(results={
        "TC-FT-001": {"status": "Fail", "detail": "boom", "evidence": None, "kind": "assert"},
    })


def test_write_prod_results(tmp_path):
    plan = _plan(tmp_path)
    out = write_prod_results(_runner(), plan, date(2026, 10, 27))
    assert out.name.endswith("_prod_smoke_2026-10-27.xlsx")
    ws = load_workbook(out)["Feature"]
    assert ws["H1"].value == "Prod Smoke Status"
    assert ws["I1"].value == "Prod Smoke Notes"
    assert ws["H2"].value == "Fail"
    assert ws["I2"].value == "boom"
    assert ws["E2"].value == "Pass"
    assert ws["H3"].value is None
    assert any("Pass,Fail,N/A,Blocking" in (dv.formula1 or "") for dv in ws.data_validations.dataValidation)
    assert load_workbook(plan)["Feature"].max_column == 7


def test_write_prod_results_twice_reuses_columns(tmp_path):
    first = write_prod_results(_runner(), _plan(tmp_path), date(2026, 10, 27))
    again = tmp_path / "again.xlsx"
    again.write_bytes(first.read_bytes())
    second = write_prod_results(_runner(), again, date(2026, 10, 28))
    ws = load_workbook(second)["Feature"]
    assert ws.max_column == 9
    assert ws["H2"].value == "Fail"


def test_allowlist_guards_prod_families_and_exact_auth_paths():
    allowlist = load_allowlist(ALLOWLIST_PATH)
    assert set(allowlist.guarded_suffixes) == {"dimagi.com", "commcarehq.org"}
    assert "/accounts/" not in allowlist.auth_path_prefixes
    assert "/o/" not in allowlist.auth_path_prefixes
