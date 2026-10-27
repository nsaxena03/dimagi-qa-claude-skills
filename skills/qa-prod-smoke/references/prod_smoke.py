"""Helpers for a prod smoke run, on top of qa-run-test-plan's harness.

    from prod_smoke import prod_guard, smoke_label, CreatedLog, write_prod_results
    from qa_runner import ConnectSession, CaseRunner, TestPlan

    guard = prod_guard()
    with ConnectSession(env="prod", guard=guard).start(org="pm_automation_01") as page:
        runner = CaseRunner(page, out_dir, guard=guard)
        ...
    write_prod_results(runner, plan)

The harness is found next to this skill, or via QA_RUN_TEST_PLAN_REFS.
"""

import os
import shutil
import sys
from datetime import date
from pathlib import Path

HERE = Path(__file__).resolve().parent
HARNESS = Path(os.getenv("QA_RUN_TEST_PLAN_REFS", HERE.parent.parent / "qa-run-test-plan" / "references"))
if str(HARNESS) not in sys.path:
    sys.path.insert(0, str(HARNESS))

from prod_guard import ProdWriteGuard, load_allowlist  # noqa: E402

ALLOWLIST_PATH = HERE / "prod_allowlist.yaml"
STATUS_HEADER = "Prod Smoke Status"
NOTES_HEADER = "Prod Smoke Notes"
STATUS_CHOICES = '"Pass,Fail,N/A,Blocking"'


def prod_guard():
    return ProdWriteGuard(load_allowlist(ALLOWLIST_PATH))


def smoke_label(ticket, today=None):
    """Name for anything a smoke run creates, so it is easy to find and clean up."""
    return f"QA-SMOKE {ticket} {(today or date.today()).strftime('%d-%b')}"


class CreatedLog:
    """What the run created on prod, so the report can list what is left behind."""

    def __init__(self):
        self._items = []

    def add(self, kind, org, name, url=None):
        self._items.append({"kind": kind, "org": org, "name": name, "url": url, "deleted": False})

    def mark_deleted(self, name):
        for item in self._items:
            if item["name"] == name:
                item["deleted"] = True

    def leftovers(self):
        return [{k: v for k, v in item.items() if k != "deleted"}
                for item in self._items if not item["deleted"]]


def _case_id(value):
    from qa_runner import _case_id as harness_case_id

    return harness_case_id(value)


def write_prod_results(runner, plan, today=None):
    """Write prod results into a dated COPY, in their own two columns.

    The plan's Status and Comments hold the staging run and stay untouched.
    `plan` may be a TestPlan or a path.
    """
    from openpyxl import load_workbook
    from openpyxl.worksheet.datavalidation import DataValidation

    from qa_runner import _header_columns

    plan_path = Path(getattr(plan, "path", plan))
    out = plan_path.with_name(
        f"{plan_path.stem}_prod_smoke_{(today or date.today()).isoformat()}{plan_path.suffix}"
    )
    shutil.copyfile(plan_path, out)

    wb = load_workbook(out)
    for ws in wb.worksheets:
        cols = _header_columns(ws)
        if "status" not in cols:
            continue
        status_col = cols.get(STATUS_HEADER.lower())
        new_columns = status_col is None
        if new_columns:
            status_col = ws.max_column + 1
            ws.cell(row=1, column=status_col, value=STATUS_HEADER)
            ws.cell(row=1, column=status_col + 1, value=NOTES_HEADER)
        notes_col = status_col + 1

        case_rows = []
        for row in range(2, ws.max_row + 1):
            case = _case_id(ws.cell(row=row, column=1).value)
            if case is None:
                continue
            case_rows.append(row)
            result = runner.results.get(case)
            if not result:
                continue
            note = result["detail"] or ""
            if result.get("evidence"):
                note += f"\nScreenshot: {result['evidence']}"
            ws.cell(row=row, column=status_col, value=result["status"])
            ws.cell(row=row, column=notes_col, value=note)

        if new_columns and case_rows:
            dv = DataValidation(type="list", formula1=STATUS_CHOICES, allow_blank=True)
            ws.add_data_validation(dv)
            letter = ws.cell(row=1, column=status_col).column_letter
            for row in case_rows:
                dv.add(f"{letter}{row}")
    wb.save(out)
    print(f"\nWrote prod smoke results into {out}")
    return out
